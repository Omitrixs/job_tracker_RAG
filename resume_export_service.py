import os
import io
from pypdf import PdfReader
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def extract_text_from_pdf(uploaded_file) -> str:
    """Extracts raw text from an uploaded PDF resume file."""
    try:
        reader = PdfReader(uploaded_file)
        text = ""
        for page in reader.pages:
            extracted = page.extract_text()
            if extracted:
                text += extracted + "\n"
        return text.strip()
    except Exception as e:
        print(f"PDF Extraction Error: {e}")
        return ""

def generate_docx_resume(optimized_text: str) -> io.BytesIO:
    """
    Parses optimized resume markdown/text and generates a clean, executive Word (.docx) document.
    """
    doc = Document()
    
    # Set standard margins (1 inch)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Configure default style font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(51, 51, 51) # Dark charcoal text

    # Process text line by line to build styled Word elements
    lines = optimized_text.split('\n')
    
    for line in lines:
        line_stripped = line.strip()
        if not line_stripped:
            continue
            
        # Heading 1 (e.g. # SUMMARY or Name)
        if line_stripped.startswith("# "):
            p = doc.add_paragraph()
            run = p.add_run(line_stripped.replace("# ", "").upper())
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 51, 102) # Executive Navy
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            
        # Heading 2 (e.g. ## Experience / Skills)
        elif line_stripped.startswith("## ") or line_stripped.startswith("### "):
            text_clean = line_stripped.replace("## ", "").replace("### ", "")
            p = doc.add_paragraph()
            run = p.add_run(text_clean.upper())
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 51, 102)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(2)
            
        # Bullet points
        elif line_stripped.startswith("- ") or line_stripped.startswith("* "):
            bullet_text = line_stripped[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            run = p.add_run(bullet_text)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            
        # Regular body paragraphs
        else:
            p = doc.add_paragraph()
            p.add_run(line_stripped)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15

    # Save to binary buffer for Streamlit download button
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer